"""
Genera js/templates.js a partire dai modelli Excel in tools/templates/*.xlsx.

Per ogni modello:
  - individua automaticamente le celle da compilare (intestazione, righe giocatori,
    falli, entrate, punti per tempo, punteggio progressivo, sospensioni, tabella punti,
    punteggio finale, istruttori);
  - incorpora il file .xlsx in base64 per l'uso offline.

Uso:  python tools/build_templates.py
I modelli si rigenerano dal file originale con tools/split_templates.ps1.
"""
import base64
import json
import os
import re

import openpyxl
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
TPL_DIR = os.path.join(HERE, "templates")
OUT = os.path.join(HERE, "..", "js", "templates.js")

KEYS = ["5c5", "4c4", "4c4open", "3c3sprint", "4c4sprint",
        "5c5_int", "4c4_int", "4c4open_int", "3c3sprint_int", "4c4sprint_int"]


def ref(r, c):
    return f"{get_column_letter(c)}{r}"


def text(ws, r, c):
    v = ws.cell(r, c).value
    return "" if v is None else str(v).strip()


def merged_at(ws, r, c):
    for m in ws.merged_cells.ranges:
        if m.min_row <= r <= m.max_row and m.min_col <= c <= m.max_col:
            return m
    return None


def extent_end(ws, r, c):
    m = merged_at(ws, r, c)
    return m.max_col if m else c


def has_bottom_border(ws, r, c):
    b = ws.cell(r, c).border
    return b is not None and b.bottom is not None and b.bottom.style is not None


def find_label(ws, pattern, rows, flags=re.I):
    rx = re.compile(pattern, flags)
    for r in rows:
        for c in range(1, ws.max_column + 1):
            if rx.search(text(ws, r, c)):
                return r, c
    return None


def value_cell_after(ws, r, c):
    """Prima cella con bordo inferiore (riga da compilare) a destra dell'etichetta."""
    start = extent_end(ws, r, c) + 1
    # l'etichetta puo' traboccare su celle vuote senza bordo: salta fino al primo bordo
    for cc in range(start, min(start + 12, ws.max_column + 1)):
        if text(ws, r, cc):
            break  # siamo arrivati a un'altra etichetta
        if has_bottom_border(ws, r, cc):
            return ref(r, cc)
    return ref(r, start)


def numeric_cells(ws, r, c0, c1):
    out = []
    for c in range(c0, c1 + 1):
        v = ws.cell(r, c).value
        if isinstance(v, (int, float)):
            out.append((int(v), c))
    return out


def analyse(key):
    path = os.path.join(TPL_DIR, key + ".xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    maxc = ws.max_column
    L = {"sheet": ws.title}

    # --- blocchi squadra -------------------------------------------------
    team_rows = []
    for r in range(1, ws.max_row + 1):
        if text(ws, r, 2).upper().startswith('SQUADRA "') and find_label(ws, r"N\.\s*Maglia", [r]):
            team_rows.append(r)
    assert len(team_rows) == 2, (key, team_rows)
    first_team = team_rows[0]

    # --- intestazione ----------------------------------------------------
    head_rows = range(1, first_team)
    header = {}
    for name, pat in [("categoria", r"^CATEGORIA"), ("girone", r"^GIRONE"), ("gara", r"^GARA N"),
                      ("data", r"^DATA"), ("ora", r"^ORA"), ("campo", r"^CAMPO DI GIOCO"),
                      ("squadraA", r'^SQUADRA "A"'), ("squadraB", r'^SQUADRA "B"'),
                      ("arbitri", r"^ARBITRI"), ("segnapunti", r"^SEGNAPUNTI"),
                      ("cronometrista", r"^CRONOMETRISTA")]:
        pos = find_label(ws, pat, head_rows)
        assert pos, (key, name)
        header[name] = value_cell_after(ws, *pos)
    L["header"] = header

    teams = []
    for h in team_rows:
        t = {"titleCell": ref(h, 2), "titleText": text(ws, h, 2),
             "colorCell": ref(h + 2, 2), "colorText": text(ws, h + 2, 2)}
        _, maglia_c = find_label(ws, r"N\.\s*Maglia", [h])
        _, entr_c = find_label(ws, r"Entrate", [h])
        _, falli_c = find_label(ws, r"^FALLI", [h])
        fm = merged_at(ws, h, falli_c)
        entr = numeric_cells(ws, h + 2, entr_c, falli_c - 1)
        # tempi: etichette "... TEMPO" nella riga h+1
        periods = []
        for c in range(1, maxc + 1):
            if re.search(r"TEMPO$", text(ws, h + 1, c)):
                m = merged_at(ws, h + 1, c)
                periods.append([m.min_col if m else c, m.max_col if m else c])
        rows = []
        r = h + 3
        while text(ws, r, 2).upper() != "ISTRUTTORE:" and r < h + 30:
            rows.append(r)
            r += 1
        coach1 = r
        t["playerRows"] = rows
        t["nameCol"] = 2
        t["numCol"] = maglia_c
        t["entrateCols"] = [c for _, c in sorted(entr)]
        t["foulCols"] = list(range(fm.min_col, fm.max_col + 1))
        t["periodCols"] = periods
        # istruttori
        coaches = []
        for cr in (coach1, coach1 + 2):
            pos = find_label(ws, r"Fallo Tec", [cr])
            lab_end = extent_end(ws, cr, pos[1])
            boxes = []
            c = lab_end + 1
            while len(boxes) < 2:
                boxes.append(ref(cr, c))
                c = extent_end(ws, cr, c) + 1
            coaches.append({"nameCell": ref(cr, 2), "nameText": text(ws, cr, 2),
                            "cardCell": ref(cr + 1, 2), "cardText": text(ws, cr + 1, 2),
                            "techCells": boxes})
        t["coaches"] = coaches
        # punteggio progressivo + sospensioni
        row1 = numeric_cells(ws, coach1, 1, maxc)
        row2 = numeric_cells(ws, coach1 + 2, 1, maxc)
        k = max(n for n, _ in row1)
        running = {}
        for n, c in row1:
            running[n] = ref(coach1, c)
        timeouts = []
        for n, c in row2:
            if n > k:
                running[n] = ref(coach1 + 2, c)
            else:
                timeouts.append((n, ref(coach1 + 2, c)))
        t["running"] = [running[i] for i in range(1, max(running) + 1)]
        t["timeouts"] = [c for _, c in sorted(timeouts)]
        teams.append(t)
    L["teams"] = teams

    # --- tabella punti per tempo ---------------------------------------
    last_coach = int(re.sub(r"\D", "", teams[1]["coaches"][1]["cardCell"]))
    ptrows = []
    for r in range(last_coach + 1, ws.max_row + 1):
        lab = text(ws, r, 6)
        if not lab:
            continue
        a = find_label(ws, r"^A$", [r])
        b = find_label(ws, r"^B$", [r])
        if not a or not b:
            continue
        av = {n: ref(r, c) for n, c in numeric_cells(ws, r, a[1] + 1, b[1] - 1)}
        bv = {n: ref(r, c) for n, c in numeric_cells(ws, r, b[1] + 1, b[1] + 8)}
        ptrows.append({"label": lab, "A": av, "B": bv})
    # righe dei tempi (in ordine) e riga dello spareggio ("Gioco" con il 2 annerito)
    L["periodPoints"] = [p for p in ptrows if re.search(r"tempo", p["label"], re.I)]
    extra = [p for p in ptrows if not re.search(r"tempo", p["label"], re.I)]
    sp = [p for p in extra if 2 not in p["A"]] or extra
    L["spareggio"] = sp[0] if sp else None
    fa = find_label(ws, r"^A[\.…]", range(last_coach + 1, ws.max_row + 1))
    fb = find_label(ws, r"^B[\.…]", range(last_coach + 1, ws.max_row + 1))
    L["final"] = {"A": ref(*fa), "B": ref(*fb)}

    # --- firme: righe (celle unite con bordo inferiore) sotto ogni etichetta ---
    labels = []
    for name in ["ISTRUTTORI", "SEGNAPUNTI", "CRONOMETRISTA", "ARBITRI"]:
        pos = find_label(ws, "^" + name, range(last_coach + 1, ws.max_row + 1))
        assert pos, (key, name)
        m = merged_at(ws, *pos)
        labels.append((name.lower(), pos[0], m.min_col, m.max_col))
    sig = {}
    for name, row, c0, c1 in labels:
        nxt = min([r for n, r, a0, a1 in labels if a0 == c0 and r > row] + [ws.max_row + 1])
        lines = sorted(m.min_row for m in ws.merged_cells.ranges
                       if m.min_col == c0 and m.max_col == c1 and row < m.min_row < nxt and has_bottom_border(ws, m.min_row, c0))
        sig[name] = [{"c0": c0, "c1": c1, "row": r} for r in lines]
    L["signatures"] = sig
    return L


def main():
    layouts, files = {}, {}
    for k in KEYS:
        layouts[k] = analyse(k)
        with open(os.path.join(TPL_DIR, k + ".xlsx"), "rb") as f:
            files[k] = base64.b64encode(f.read()).decode()
    with open(OUT, "w", encoding="utf-8") as f:
        f.write("/* File generato da tools/build_templates.py - non modificare a mano */\n")
        f.write("window.REFERTO_LAYOUTS = " + json.dumps(layouts, ensure_ascii=False) + ";\n")
        f.write("window.REFERTO_FILES = " + json.dumps(files) + ";\n")
    for k in KEYS:
        print(k, "firme", {n: [x["row"] for x in v] for n, v in layouts[k]["signatures"].items()})
    for k in KEYS:
        L = layouts[k]
        print(k, L["header"], "rows", len(L["teams"][0]["playerRows"]), "entr", len(L["teams"][0]["entrateCols"]),
              "periods", len(L["teams"][0]["periodCols"]), "running", len(L["teams"][0]["running"]),
              "TO", len(L["teams"][0]["timeouts"]), "pt", [p["label"] for p in L["periodPoints"]], L["final"])


if __name__ == "__main__":
    main()
